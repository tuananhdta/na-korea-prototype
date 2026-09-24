import zipfile
import os
from pptx import Presentation

pptx_path = "/Users/tuananh/Na-Korea/na-korea-prototype/2 file mẫu cho catalog/Đào tạo nội bộ - Ginsenoside .pptx"

prs = Presentation(pptx_path)
print(f"Total slides in PPTX: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width.inches} inches, height: {prs.slide_height.inches} inches")

for i, slide in enumerate(prs.slides):
    shapes = [shape.name for shape in slide.shapes]
    text = " ".join([shape.text for shape in slide.shapes if shape.has_text_frame])
    print(f"Slide {i+1} ({len(shapes)} shapes): {text[:60]}...")
